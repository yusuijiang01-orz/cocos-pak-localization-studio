#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Apply a translated term glossary workbook to an exported localization XLSX.

This is intentionally deterministic and offline: it never calls an LLM.  It
rewrites only the editable `text` column of the selected exported workbook, so
existing Studio import/build validators remain the authority for PAK safety.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
import traceback
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
VI_RE = re.compile(
    r"[ÀÁÂÃÈÉÊÌÍÒÓÔÕÙÚÝàáâãèéêìíòóôõùúý"
    r"ĂăĐđĨĩŨũƠơƯưẠ-ỹ]"
)
PAIR_SPLIT_RE = re.compile(r"\s*(?:=>|=|：|:)\s*", re.UNICODE)


def emit(obj: dict[str, Any]) -> None:
    print(json.dumps(obj, ensure_ascii=False), flush=True)


def text_value(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value)).strip()


def norm_header(value: Any) -> str:
    return re.sub(r"[\s_\-]+", "", text_value(value).casefold())


def norm_key(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value or "")).strip()).casefold()


def contains_cjk(value: str) -> bool:
    return bool(CJK_RE.search(value or ""))


def load_json(path: Path | None) -> dict[str, Any]:
    if not path or not str(path) or not Path(path).is_file():
        return {}
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def column_index(headers: list[Any], names: set[str], fallback: int | None = None) -> int | None:
    normalized = [norm_header(h) for h in headers]
    for index, header in enumerate(normalized, 1):
        if header in names:
            return index
    return fallback


def worksheet_headers(ws) -> list[Any]:
    return [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]


def mapping_rows(mapping: dict[str, Any]) -> list[list[Any]]:
    rows = mapping.get("rows")
    return rows if isinstance(rows, list) else []


def find_text_column(headers: list[Any]) -> int:
    idx = column_index(headers, {"text", "译文", "中文", "translation", "translated", "target", "textzh", "textcn"})
    if idx:
        return idx
    return max(1, len(headers))


def find_id_column(headers: list[Any]) -> int:
    return column_index(headers, {"id", "编号", "序号"}, fallback=1) or 1


def load_glossary_from_pairs(ws, headers: list[Any]) -> dict[str, str]:
    source_col = column_index(headers, {"source", "src", "原文", "源文", "术语", "term", "text", "viet", "vi", "越南文"}, fallback=1) or 1
    target_col = column_index(headers, {"target", "dst", "translation", "translated", "zh", "cn", "中文", "译文", "textzh", "textcn"})
    if not target_col and ws.max_column >= 2:
        target_col = 2 if source_col != 2 else 1
    pairs: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2):
        source = text_value(row[source_col - 1].value) if source_col <= len(row) else ""
        target = text_value(row[target_col - 1].value) if target_col and target_col <= len(row) else ""
        if not target and source:
            parts = PAIR_SPLIT_RE.split(source, maxsplit=1)
            if len(parts) == 2:
                source, target = text_value(parts[0]), text_value(parts[1])
        if source and target and norm_key(source) != norm_key(target) and contains_cjk(target):
            pairs[source] = target
    return pairs


def load_glossary_from_mapped_single_column(ws, mapping: dict[str, Any], headers: list[Any]) -> dict[str, str]:
    rows = mapping_rows(mapping)
    text_col = find_text_column(headers)
    pairs: dict[str, str] = {}
    for idx, row in enumerate(ws.iter_rows(min_row=2), 0):
        if idx >= len(rows):
            break
        mapped = rows[idx]
        source = text_value(mapped[1] if len(mapped) > 1 else "")
        target = text_value(row[text_col - 1].value) if text_col <= len(row) else ""
        if source and target and norm_key(source) != norm_key(target) and contains_cjk(target):
            pairs[source] = target
    return pairs


def load_glossary(glossary_xlsx: Path, glossary_mapping: Path | None) -> tuple[dict[str, str], dict[str, Any]]:
    mapping = load_json(glossary_mapping)
    wb = load_workbook(glossary_xlsx)
    ws = wb.active
    headers = worksheet_headers(ws)
    pairs = load_glossary_from_pairs(ws, headers)
    mapped_pairs: dict[str, str] = {}
    if mapping.get("mode") == "multi-pak-safe-term-glossary":
        mapped_pairs = load_glossary_from_mapped_single_column(ws, mapping, headers)
    # Explicit two-column pairs win; mapped single-column pairs make the current
    # exported one-column glossary usable after Google/Ollama translation.
    merged = {**mapped_pairs, **pairs}
    by_key: dict[str, str] = {}
    deduped: dict[str, str] = {}
    for source, target in sorted(merged.items(), key=lambda item: len(item[0]), reverse=True):
        key = norm_key(source)
        if not key or key in by_key:
            continue
        by_key[key] = target
        deduped[source] = target
    report = {
        "glossary_rows": max(0, ws.max_row - 1),
        "glossary_terms": len(deduped),
        "glossary_mapping": str(glossary_mapping or ""),
    }
    return deduped, report


def compile_terms(pairs: dict[str, str]) -> list[tuple[re.Pattern[str], str, str]]:
    terms: list[tuple[re.Pattern[str], str, str]] = []
    for source, target in sorted(pairs.items(), key=lambda item: len(item[0]), reverse=True):
        source = text_value(source)
        target = text_value(target)
        if not source or not target:
            continue
        pattern = re.compile(re.escape(source), re.IGNORECASE)
        terms.append((pattern, source, target))
    return terms


def translate_cell(value: str, terms: list[tuple[re.Pattern[str], str, str]]) -> tuple[str, int, list[str]]:
    current = text_value(value)
    if not current:
        return current, 0, []
    applied: list[str] = []
    count = 0
    # Fast exact match before substring replacement.
    for _pattern, source, target in terms:
        if norm_key(current) == norm_key(source):
            return target, 1, [source]
    for pattern, source, target in terms:
        current, replacements = pattern.subn(target, current)
        if replacements:
            count += replacements
            applied.append(source)
    return current, count, applied


def apply_glossary(source_xlsx: Path, source_mapping: Path | None,
                   glossary_xlsx: Path, glossary_mapping: Path | None,
                   output_xlsx: Path) -> dict[str, Any]:
    if not source_xlsx.is_file():
        raise FileNotFoundError(f"导出 XLSX 不存在：{source_xlsx}")
    if not glossary_xlsx.is_file():
        raise FileNotFoundError(f"术语库 XLSX 不存在：{glossary_xlsx}")

    glossary_pairs, glossary_report = load_glossary(glossary_xlsx, glossary_mapping)
    if not glossary_pairs:
        raise ValueError("术语库没有可用的 原文 → 中文 术语对。请确认术语库已翻译，或添加第二列中文译文。")
    terms = compile_terms(glossary_pairs)

    mapping = load_json(source_mapping)
    mapped_rows = mapping_rows(mapping)
    mapped_by_id = {
        text_value(row[0]): text_value(row[2] if len(row) > 2 else row[1])
        for row in mapped_rows
        if isinstance(row, list) and row
    }

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    if output_xlsx.resolve() != source_xlsx.resolve():
        shutil.copy2(source_xlsx, output_xlsx)
    wb = load_workbook(output_xlsx)
    ws = wb.active
    headers = worksheet_headers(ws)
    id_col = find_id_column(headers)
    text_col = find_text_column(headers)

    total_rows = changed_rows = unchanged_rows = blank_rows = 0
    matched_terms_total = 0
    examples: list[dict[str, Any]] = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2), 2):
        total_rows += 1
        id_value = text_value(row[id_col - 1].value) if id_col <= len(row) else ""
        text_cell = row[text_col - 1]
        current_text = text_value(text_cell.value)
        source_text = current_text or mapped_by_id.get(id_value, "")
        if not source_text:
            blank_rows += 1
            continue
        translated, replacements, applied_terms = translate_cell(source_text, terms)
        if replacements and translated != current_text:
            text_cell.value = translated
            changed_rows += 1
            matched_terms_total += replacements
            if len(examples) < 30:
                examples.append({
                    "row": row_number,
                    "id": id_value,
                    "before": source_text,
                    "after": translated,
                    "terms": applied_terms[:8],
                })
        else:
            unchanged_rows += 1
        if total_rows % 1000 == 0:
            emit({
                "event": "progress",
                "phase": "glossary-translate",
                "percent": min(90, 10 + round(total_rows / max(1, ws.max_row - 1) * 80, 1)),
                "message": f"正在应用术语库：{total_rows:,}/{max(0, ws.max_row - 1):,}",
            })
    wb.save(output_xlsx)
    return {
        "mode": "glossary-xlsx-translate",
        "source_xlsx": str(source_xlsx.resolve()),
        "source_mapping": str(source_mapping or ""),
        "glossary_xlsx": str(glossary_xlsx.resolve()),
        "output_xlsx": str(output_xlsx.resolve()),
        "mapping_mode": mapping.get("mode", ""),
        "mapping_version": mapping.get("version", ""),
        "paks": mapping.get("paks") or sorted({item[1] for item in mapping.get("records", []) if isinstance(item, list) and len(item) > 1}),
        "workbook_rows": total_rows,
        "translated_rows": changed_rows,
        "unchanged_rows": unchanged_rows,
        "blank_rows": blank_rows,
        "matched_terms": matched_terms_total,
        "examples": examples,
        **glossary_report,
    }


def main(argv: list[str]) -> int:
    if len(argv) != 6:
        emit({"event": "error", "message": "Usage: glossary_xlsx_translate.py <source_xlsx> <source_mapping_json> <glossary_xlsx> <glossary_mapping_json_or_empty> <output_xlsx>"})
        return 2
    try:
        source_xlsx = Path(argv[1])
        source_mapping = Path(argv[2]) if argv[2] else None
        glossary_xlsx = Path(argv[3])
        glossary_mapping = Path(argv[4]) if argv[4] else None
        output_xlsx = Path(argv[5])
        emit({"event": "progress", "phase": "glossary-translate", "percent": 3, "message": "正在读取导出 XLSX 和术语库 XLSX…"})
        report = apply_glossary(source_xlsx, source_mapping, glossary_xlsx, glossary_mapping, output_xlsx)
        emit({"event": "done", "report": report})
        return 0
    except Exception as exc:
        emit({"event": "error", "message": str(exc), "trace": traceback.format_exc()})
        return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
